"""Deck-level behavior: multi-slide assembly and cross-slide lint."""
import json
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

import build_deck
import lint_deck


def _slide_spec(title_text, font="Arial", size=30, role="heading",
                color="112233", x=500000):
    return {
        "spec_version": 1,
        "slide": {"width_emu": 12192000, "height_emu": 6858000},
        "elements": [
            {"id": "h", "type": "textbox", "role": role,
             "box": {"x": x, "y": 500000, "cx": 6000000, "cy": 800000},
             "paragraphs": [{"lines": [{"text": title_text, "font": font,
                                        "size_pt": size, "color": color,
                                        "bold": True}]}]},
        ],
    }


@pytest.fixture()
def project(tmp_path):
    (tmp_path / "specs").mkdir()
    return tmp_path


def _write_deck(project, slide_specs, tokens=None):
    paths = []
    for i, spec in enumerate(slide_specs):
        p = project / "specs" / f"s{i}.spec.json"
        p.write_text(json.dumps(spec))
        paths.append(f"specs/s{i}.spec.json")
    deck = {
        "deck_version": 1,
        "meta": {"doc_props": {"title": "Deck", "author": "A"}},
        "slide": {"width_emu": 12192000, "height_emu": 6858000},
        "slides": paths,
    }
    if tokens:
        deck["tokens"] = tokens
    deck_path = project / "specs" / "t.deck.json"
    deck_path.write_text(json.dumps(deck))
    return deck_path, deck


def test_deck_assembles_all_slides_in_order(project):
    deck_path, deck = _write_deck(project, [
        _slide_spec("First"), _slide_spec("Second"), _slide_spec("Third")])
    errors, slide_specs = build_deck.validate_deck(deck, project)
    assert errors == []
    prs = build_deck.build_deck(deck, slide_specs, project_root=project)
    out = project / "deck.pptx"
    prs.save(out)
    assert len(Presentation(out).slides) == 3
    with zipfile.ZipFile(out) as z:
        assert "<Slides>3</Slides>" in z.read("docProps/app.xml").decode()
    texts = ["".join(r.text for s in [sl] for sh in sl.shapes
             if sh.has_text_frame for p in sh.text_frame.paragraphs
             for r in p.runs) for sl in Presentation(out).slides]
    assert texts == ["First", "Second", "Third"]


def test_deck_rejects_mismatched_slide_size(project):
    bad = _slide_spec("Odd one")
    bad["slide"] = {"width_emu": 9144000, "height_emu": 6858000}
    deck_path, deck = _write_deck(project, [_slide_spec("ok"), bad])
    errors, _ = build_deck.validate_deck(deck, project)
    assert any("deck demands" in e for e in errors)


def test_lint_deck_font_budget_and_type_scale(project):
    tokens = {
        "max_font_families": 2,
        "type_scale": {"heading": {"font": "Arial", "size_pt": 30, "bold": True}},
    }
    deck_path, _ = _write_deck(project, [
        _slide_spec("A", font="Arial"),
        _slide_spec("B", font="Georgia", size=28),   # off-token role
        _slide_spec("C", font="Verdana"),            # third family
    ], tokens=tokens)
    report = lint_deck.lint_deck(deck_path, project_root=project)
    checks = {(f["check"], f["severity"]) for f in report["findings"]}
    assert ("font_budget", "error") in checks
    assert ("type_scale", "error") in checks


def test_lint_deck_role_alignment_warns_on_grid_drift(project):
    deck_path, _ = _write_deck(project, [
        _slide_spec("A", x=500000),
        _slide_spec("B", x=1200000),   # same role, different grid line
    ])
    report = lint_deck.lint_deck(deck_path, project_root=project)
    assert any(f["check"] == "role_alignment" for f in report["findings"])
    assert report["n_errors"] == 0


def test_lint_deck_clean_deck_passes(project):
    tokens = {
        "palette_roles": {"ink": "112233"},
        "type_scale": {"heading": {"font": "Arial", "size_pt": 30, "bold": True}},
        "margins_emu": {"left": 400000, "right": 400000},
        "max_font_families": 2,
    }
    deck_path, _ = _write_deck(project, [
        _slide_spec("A"), _slide_spec("B")], tokens=tokens)
    report = lint_deck.lint_deck(deck_path, project_root=project)
    assert report["n_errors"] == 0 and report["n_warns"] == 0
