"""Behavioral tests for native_packaging.apply_packaging: a built deck
must carry the event's theme in its actual theme part (what PowerPoint's
pickers read) and no unused stock layouts — while still being a valid,
reopenable package.
"""
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

import build_deck
from native_packaging import apply_packaging

SYNTH = {
    "theme_source": {
        "kind": "synthesized",
        "name": "Dental Test",
        "colors": {"dk1": "291A12", "lt1": "FAF5E8", "dk2": "493D3A",
                   "lt2": "E8E0D2", "accent1": "6A5337", "accent2": "A68A69",
                   "accent3": "4C4339", "accent4": "B1926B",
                   "accent5": "635444", "accent6": "AFA59A"},
        "major_font": "Noto Serif",
        "minor_font": "Roboto Condensed",
    },
    "strip_unused_layouts": True,
}

MINIMAL_SPEC = {
    "spec_version": 1,
    "slide": {"width_emu": 12192000, "height_emu": 6858000},
    "elements": [
        {"id": "t", "type": "textbox",
         "box": {"x": 0, "y": 0, "cx": 1000000, "cy": 500000},
         "paragraphs": [{"lines": [{"text": "x", "font": "Arial",
                                    "size_pt": 12, "color": "000000"}]}]},
    ],
}


def _build_to(path, packaging):
    spec = dict(MINIMAL_SPEC, packaging=packaging)
    prs = build_deck.build(spec)
    prs.save(path)
    apply_packaging(path, packaging, Path("/nonexistent-root-unused"))
    return path


@pytest.fixture(scope="module")
def synth_pptx(tmp_path_factory):
    return _build_to(tmp_path_factory.mktemp("pkg") / "synth.pptx", SYNTH)


def test_synthesized_theme_reaches_theme_part(synth_pptx):
    with zipfile.ZipFile(synth_pptx) as z:
        theme = z.read("ppt/theme/theme1.xml").decode()
    assert 'val="6A5337"' in theme        # accent1
    assert 'val="FAF5E8"' in theme        # lt1
    assert 'typeface="Noto Serif"' in theme
    assert 'typeface="Roboto Condensed"' in theme
    assert 'name="Dental Test"' in theme
    assert "4F81BD" not in theme          # stock Office accent1 must be gone


def test_unused_stock_layouts_stripped(synth_pptx):
    with zipfile.ZipFile(synth_pptx) as z:
        names = z.namelist()
        layouts = [n for n in names
                   if n.startswith("ppt/slideLayouts/slideLayout")]
        master = z.read("ppt/slideMasters/slideMaster1.xml").decode()
        cts = z.read("[Content_Types].xml").decode()
    assert len(layouts) == 1, f"expected only the used layout, got {layouts}"
    assert master.count("<p:sldLayoutId ") == 1
    assert cts.count('PartName="/ppt/slideLayouts/') == 1


def test_package_still_opens_and_renders_layout_link(synth_pptx):
    prs = Presentation(synth_pptx)
    assert len(prs.slides) == 1
    # the surviving layout is still resolvable from the slide
    assert prs.slides[0].slide_layout is not None


def test_template_theme_graft(tmp_path):
    # forge an "organizer template": stock package with a recognizable theme
    donor = tmp_path / "donor_template.pptx"
    Presentation().save(donor)
    with zipfile.ZipFile(donor) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["ppt/theme/theme1.xml"] = (
        entries["ppt/theme/theme1.xml"]
        .replace(b'val="4F81BD"', b'val="12AB34"')
        .replace(b'name="Office"', b'name="GraftBrand"', 1))
    with zipfile.ZipFile(donor, "w") as z:
        for n, b in entries.items():
            z.writestr(n, b)

    packaging = {"theme_source": {"kind": "template",
                                  "path": "donor_template.pptx",
                                  "theme_index": 0},
                 "strip_unused_layouts": True}
    out = tmp_path / "grafted.pptx"
    spec = dict(MINIMAL_SPEC, packaging=packaging)
    prs = build_deck.build(spec)
    prs.save(out)
    apply_packaging(out, packaging, tmp_path)

    with zipfile.ZipFile(out) as z:
        theme = z.read("ppt/theme/theme1.xml").decode()
    assert 'val="12AB34"' in theme
    assert 'name="GraftBrand"' in theme
    assert Presentation(out).slides is not None
