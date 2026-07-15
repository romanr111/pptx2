import shutil
import subprocess
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import render


def _docker_is_usable():
    if shutil.which("docker") is None:
        return False
    try:
        return subprocess.run(
            ["docker", "info"], capture_output=True, timeout=5,
        ).returncode == 0
    except OSError:
        return False
    except subprocess.TimeoutExpired:
        return False


@pytest.mark.docker
@pytest.mark.skipif(not _docker_is_usable(), reason="Docker daemon is unavailable")
def test_docker_renderer_converts_a_two_slide_deck_to_two_pngs(tmp_path):
    deck = tmp_path / "two-slide.pptx"
    shared_logo = tmp_path / "shared-logo.png"
    Image.new("RGBA", (320, 96), (91, 73, 211, 255)).save(shared_logo)
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(
            str(shared_logo), Inches(10.2), Inches(0.35), width=Inches(2.0)
        )
    presentation.save(deck)

    with zipfile.ZipFile(deck) as package:
        shared_media = [
            name for name in package.namelist() if name.startswith("ppt/media/")
        ]
    assert len(shared_media) == 1

    output = tmp_path / "rendered"
    pngs = render.render_all(deck, output, renderer="docker")

    assert (output / "two-slide.pdf").stat().st_size > 0
    assert len(pngs) == 2
    assert all(path.stat().st_size > 0 for path in pngs)
